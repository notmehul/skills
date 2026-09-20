"""Render workspace/<slug>/sheet/*.md into an xlsx workbook with native charts.

Run: python3 render_sheet.py workspace/<slug> [--out path.xlsx] [--skip-checks]

Runs check_sheet first and writes checks.md. Renders even when checks fail so
the analyst can see the sheet, but exits 1 so the caller knows to fix and
re-render. Requires openpyxl (scripts/setup.sh installs it into .venv).
"""

from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import check_sheet  # noqa: E402
from sheet_md import (  # noqa: E402
    DIRECTION_RE, INLINE_LINK_RE, LINK_RE, STATUS_RE, STATUS_WORDS, Chart, Tab, Table, cell_kind, parse_tab,
)

# ---- style -----------------------------------------------------------------
FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=15, bold=True)
F_SUB = Font(name=FONT, size=9, color="6B6B6B")
F_SECTION = Font(name=FONT, size=12, bold=True)
F_NOTE = Font(name=FONT, size=9, color="6B6B6B")
F_HEADER = Font(name=FONT, size=9, bold=True, color="333333")
F_LINK = Font(name=FONT, size=10, color="1155CC", underline="single")
F_GLYPH = Font(name=FONT, size=11, color="333333")
F_GLYPH_HL = Font(name=FONT, size=11, color="1F6F8B", bold=True)
F_GREEN = Font(name=FONT, size=10, color="2E7D32", bold=True)
F_RED = Font(name=FONT, size=10, color="C62828", bold=True)
F_AMBER = Font(name=FONT, size=10, color="B26A00", bold=True)
F_GREY = Font(name=FONT, size=10, color="8A8A8A")

ACCENT = "1F6F8B"
BAR_GREY = "B8B8B8"
FILL_HL = PatternFill("solid", fgColor="E8F1F5")
RULE = Side(style="thin", color="8C8C8C")
HAIR = Side(style="thin", color="E6E6E6")
B_HEADER = Border(bottom=RULE)
B_ROW = Border(bottom=HAIR)

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
RIGHT = Alignment(wrap_text=True, vertical="top", horizontal="right")
ROTATE = Alignment(textRotation=90, vertical="bottom", horizontal="center")

MARGIN_COL = 1          # column A stays empty
FIRST_COL = 2           # content starts at B
MIN_CONTENT_COLS = 5
CHART_WIDTH_CM = 15.0
ROW_CM = 0.53           # default row height in cm, for chart stacking


def clamp(v, lo, hi):
    return max(lo, min(hi, v))


# ---- widths ------------------------------------------------------------------
def compute_widths(tab: Tab) -> list[float]:
    """Per content-column widths in Excel units. Tables in a tab share columns."""
    if tab.widths:
        return list(tab.widths)
    widths: dict[int, float] = {}
    for kind, payload in tab.blocks:
        if kind != "table":
            continue
        t: Table = payload
        if t.matrix:
            cand = [30.0] + [6.5] * (t.ncols - 2) + [40.0] if t.ncols > 2 else [30.0, 40.0]
        elif t.kv:
            cand = [24.0, 70.0]
        else:
            cand = []
            for j in range(t.ncols):
                lens = [len(r[j]) for r in t.data_rows() if j < len(r)]
                lens.sort()
                p90 = lens[int(0.9 * (len(lens) - 1))] if lens else 0
                base = max(len(t.headers[j]), p90)
                # long text columns wrap; cap so rows stay readable
                w = clamp(base * 0.95, 8, 56)
                if cell_kind(t.rows[0][j] if t.rows and isinstance(t.rows[0], list) else "", t.headers[j]) in ("glyph", "center", "num"):
                    w = clamp(base * 0.9, 8, 14)
                cand.append(w)
        for j, w in enumerate(cand):
            widths[j] = max(widths.get(j, 0), w)
    n = max(len(widths), MIN_CONTENT_COLS)
    out = [widths.get(j, 12.0) for j in range(n)]
    return out


# ---- renderer ------------------------------------------------------------------
class TabRenderer:
    def __init__(self, ws, tab: Tab, chart_ws, chart_state: dict):
        self.ws = ws
        self.tab = tab
        self.chart_ws = chart_ws
        self.chart_state = chart_state  # {"row": next free row on chart data sheet}
        self.widths = compute_widths(tab)
        self.ncols = len(self.widths)
        self.r = 2
        self.chart_cursor = 2
        ws.sheet_view.showGridLines = False
        ws.column_dimensions[get_column_letter(MARGIN_COL)].width = 2
        for j, w in enumerate(self.widths):
            ws.column_dimensions[get_column_letter(FIRST_COL + j)].width = w
        ws.row_dimensions[1].height = 8

    # -- primitives
    def _cell(self, r, c, value=None, font=F_BODY, align=WRAP, fill=None, border=None, fmt=None):
        cell = self.ws.cell(row=r, column=c)
        if value is not None:
            cell.value = value
        cell.font = font
        cell.alignment = align
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
        if fmt:
            cell.number_format = fmt
        return cell

    def _merged_text(self, text, font, c1=FIRST_COL, c2=None, border=None):
        c2 = c2 or (FIRST_COL + self.ncols - 1)
        self._cell(self.r, c1, text, font, WRAP, border=border)
        if c2 > c1:
            self.ws.merge_cells(start_row=self.r, start_column=c1, end_row=self.r, end_column=c2)
            if border is not None:
                for c in range(c1 + 1, c2 + 1):
                    self.ws.cell(row=self.r, column=c).border = border
        self._autoheight(self.r, text, c1, c2, font.size or 10)
        self.r += 1

    def _autoheight(self, r, text, c1, c2, size):
        w = sum(self.widths[c1 - FIRST_COL:c2 - FIRST_COL + 1])
        per_line = max(20, int(w * 1.15 * 10 / size))
        lines = max(1, math.ceil(len(text) / per_line))
        self.ws.row_dimensions[r].height = lines * (size * 1.35) + 4

    # -- blocks
    def title(self):
        self._merged_text(self.tab.title, F_TITLE)
        if self.tab.subtitle:
            self._merged_text(self.tab.subtitle, F_SUB)
        self.r += 1

    def section(self, text):
        self.r += 0 if self.r == 2 else 1
        self._merged_text(text, F_SECTION)

    def note(self, text):
        self._merged_text(text, F_NOTE)

    def para(self, text):
        self._merged_text(text, F_BODY)

    def bullet(self, text):
        self._merged_text("• " + text, F_BODY)

    def table(self, t: Table):
        if t.kv:
            return self._kv(t)
        hr = self.r
        for j, h in enumerate(t.headers):
            c = FIRST_COL + j
            kind = cell_kind("", h)
            if t.matrix and j not in (0, t.ncols - 1):
                al = ROTATE
            elif kind in ("center", "num", "glyph"):
                al = CENTER
            else:
                al = WRAP
            self._cell(hr, c, h, F_HEADER, al, fill=FILL_HL if j in t.highlight_cols else None, border=B_HEADER)
        if t.matrix:
            self.ws.row_dimensions[hr].height = 110
        self.r += 1
        for row in t.rows:
            if isinstance(row, tuple):
                self._merged_text(row[1], F_BOLD, FIRST_COL, FIRST_COL + t.ncols - 1, border=B_ROW)
                continue
            for j, raw in enumerate(row):
                self._data_cell(self.r, FIRST_COL + j, raw, t.headers[j], j in t.highlight_cols, t.matrix)
            self.r += 1
        self.r += 1

    def _kv(self, t: Table):
        for row in t.rows:
            if isinstance(row, tuple):
                self._merged_text(row[1], F_BOLD, border=B_ROW)
                continue
            k, v = row[0], row[1]
            self._cell(self.r, FIRST_COL, k, F_BOLD, WRAP, border=B_ROW)
            c2 = FIRST_COL + self.ncols - 1
            self._cell(self.r, FIRST_COL + 1, self._plain(v), F_BODY, WRAP, border=B_ROW)
            self._apply_link(self.r, FIRST_COL + 1, v)
            if c2 > FIRST_COL + 1:
                self.ws.merge_cells(start_row=self.r, start_column=FIRST_COL + 1, end_row=self.r, end_column=c2)
                for c in range(FIRST_COL + 2, c2 + 1):
                    self.ws.cell(row=self.r, column=c).border = B_ROW
            self._autoheight(self.r, v, FIRST_COL + 1, c2, 10)
            self.r += 1
        self.r += 1

    @staticmethod
    def _plain(text: str) -> str:
        m = LINK_RE.match(text)
        if m:
            return m.group(1)
        return INLINE_LINK_RE.sub(r"\1", text)

    def _apply_link(self, r, c, raw: str):
        m = LINK_RE.match(raw) or INLINE_LINK_RE.search(raw)
        if m:
            cell = self.ws.cell(row=r, column=c)
            cell.hyperlink = m.group(2)
            cell.font = F_LINK

    def _data_cell(self, r, c, raw: str, header: str, highlight: bool, matrix: bool):
        kind = cell_kind(raw, header)
        font, align, fill, fmt, value = F_BODY, WRAP, (FILL_HL if highlight else None), None, self._plain(raw)
        if raw == "":
            self._cell(r, c, None, F_BODY, WRAP, fill=fill, border=B_ROW)
            return
        if kind == "glyph":
            font, align = (F_GLYPH_HL if highlight else F_GLYPH), CENTER
        elif kind == "direction":
            m = DIRECTION_RE.match(raw)
            font, align = (F_GREEN if m.group(1) == "▲" else F_RED), CENTER
        elif kind == "status":
            sym = raw[0] if STATUS_RE.match(raw) else STATUS_WORDS.get(raw.strip().lower(), "?")
            font = {"✔": F_GREEN, "◐": F_AMBER, "✘": F_RED}.get(sym, F_GREY)
            align = CENTER
        elif kind == "num":
            align = RIGHT
            try:
                value = float(raw.replace(",", ""))
                fmt = "#,##0.0" if value != int(value) else "#,##0"
            except ValueError:
                value = raw
        elif kind == "center":
            align = CENTER
            if re.fullmatch(r"\d+", raw):
                value = int(raw)
        if highlight and kind in ("text", "link"):
            font = F_BOLD
        self._cell(r, c, value, font, align, fill=fill, border=B_ROW, fmt=fmt)
        if kind == "link" or INLINE_LINK_RE.search(raw):
            self._apply_link(r, c, raw)

    def chart(self, ch: Chart):
        ws_d = self.chart_ws
        start = self.chart_state["row"]
        ws_d.cell(row=start, column=1, value=f"{self.tab.name}: {ch.title}").font = F_BOLD
        hdr = start + 1
        nser = len(ch.rows[0][1])
        ws_d.cell(row=hdr, column=1, value="label").font = F_HEADER
        for j in range(nser):
            name = ch.series[j] if j < len(ch.series) else ("value" if nser == 1 else f"series {j + 1}")
            ws_d.cell(row=hdr, column=2 + j, value=name).font = F_HEADER
        first = hdr + 1
        for i, (label, vals) in enumerate(ch.rows):
            ws_d.cell(row=first + i, column=1, value=label).font = F_BODY
            for j, v in enumerate(vals):
                c = ws_d.cell(row=first + i, column=2 + j, value=round(v, 2))
                c.font = F_BODY
        last = first + len(ch.rows) - 1
        self.chart_state["row"] = last + 2

        bar = BarChart()
        bar.type = "bar"
        bar.style = 10
        bar.title = ch.title
        bar.legend = None
        if nser > 1:
            bar.grouping = "stacked"
            bar.overlap = 100
            bar.legend = bar.legend or None
        n = len(ch.rows)
        bar.height = clamp(0.9 * n + 1.8, 4.5, 12)
        bar.width = CHART_WIDTH_CM
        data = Reference(ws_d, min_col=2, max_col=1 + nser, min_row=hdr, max_row=last)
        cats = Reference(ws_d, min_col=1, min_row=first, max_row=last)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.y_axis.majorGridlines = None
        bar.y_axis.delete = True
        bar.x_axis.delete = False
        bar.x_axis.scaling.orientation = "maxMin"
        bar.gapWidth = 60
        colors = [BAR_GREY, "C62828"] if nser == 2 else [BAR_GREY]
        if nser == 2 and ch.series and ch.series[0].lower().startswith("tail"):
            colors = ["2E7D32", "C62828"]
        for si, s in enumerate(bar.series):
            col = colors[si % len(colors)]
            s.graphicalProperties.solidFill = col
            s.graphicalProperties.line.solidFill = col
            s.dLbls = DataLabelList()
            s.dLbls.showVal = True
            s.dLbls.showSerName = False
            s.dLbls.showCatName = False
            s.dLbls.showLegendKey = False
            s.dLbls.numFmt = ch.fmt
            if ch.highlight is not None and si == 0:
                pt = DataPoint(idx=ch.highlight)
                pt.graphicalProperties = GraphicalProperties(solidFill=ACCENT)
                pt.graphicalProperties.line.solidFill = ACCENT
                s.dPt.append(pt)
        if nser > 1:
            from openpyxl.chart.legend import Legend
            bar.legend = Legend()
            bar.legend.position = "b"
        anchor_row = max(self.r, self.chart_cursor)
        anchor_col = FIRST_COL + self.ncols + 1
        self.ws.add_chart(bar, f"{get_column_letter(anchor_col)}{anchor_row}")
        self.chart_cursor = anchor_row + math.ceil(bar.height / ROW_CM) + 1

    def build(self):
        self.title()
        for kind, payload in self.tab.blocks:
            getattr(self, kind)(payload)


def sheet_name(title: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", title).strip()[:31] or "Sheet"
    name, n = base, 2
    while name in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def render(workspace: Path, out: Path | None = None) -> Path:
    tabs = []
    for p in sorted((workspace / "sheet").glob("*.md")):
        try:
            tabs.append(parse_tab(p.read_text(encoding="utf-8"), p))
        except ValueError as exc:
            print(f"skipped {p.name}: {exc}")
    if not tabs:
        raise SystemExit(f"no renderable tab files in {workspace / 'sheet'}")
    wb = Workbook()
    chart_ws = wb.active
    chart_ws.title = "chart data"
    chart_ws.column_dimensions["A"].width = 40
    chart_ws["A1"] = "Series behind the charts. Edit a value here and the chart updates."
    chart_ws["A1"].font = F_NOTE
    chart_state = {"row": 3}
    used: set[str] = {"chart data"}
    for tab in tabs:
        ws = wb.create_sheet(sheet_name(tab.name, used))
        TabRenderer(ws, tab, chart_ws, chart_state).build()
    wb.move_sheet("chart data", offset=len(wb.sheetnames))
    if out is None:
        company = _company_name(workspace)
        out = workspace / "output" / f"{company}_MI.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _company_name(workspace: Path) -> str:
    brief = workspace / "brief.md"
    if brief.exists():
        for line in brief.read_text(encoding="utf-8").splitlines():
            if line.startswith("# "):
                return re.sub(r"[^\w\- ]", "", line[2:]).strip().replace(" ", "_")[:40] or "Deal"
    return workspace.name


def main():
    ap = argparse.ArgumentParser(description="Render sheet/*.md into an xlsx workbook.")
    ap.add_argument("workspace")
    ap.add_argument("--out")
    ap.add_argument("--skip-checks", action="store_true")
    args = ap.parse_args()
    ws = Path(args.workspace)
    rc = 0 if args.skip_checks else check_sheet.run(ws)
    out = render(ws, Path(args.out) if args.out else None)
    print(f"rendered {out}")
    if rc:
        print("checks failed: fix the items in checks.md and render again")
    sys.exit(rc)


if __name__ == "__main__":
    main()
